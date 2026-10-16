import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_DARK = RGBColor(6, 6, 18)
    CARD_BG = RGBColor(13, 13, 43)
    ACCENT_TEAL = RGBColor(0, 229, 160)
    ACCENT_CYAN = RGBColor(0, 180, 216)
    ACCENT_RED = RGBColor(255, 71, 87)
    TEXT_LIGHT = RGBColor(238, 238, 245)
    TEXT_MUTED = RGBColor(136, 136, 170)
    WHITE = RGBColor(255, 255, 255)

    blank_layout = prs.slide_layouts[6]

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="PRICEGUARD AI • AIML-08"):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_TEAL
        p_cat.font.name = "Arial"

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT
        p_title.font.name = "Arial"

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_background(slide1)

    # Accent Glow box
    add_card(slide1, 1.0, 1.2, 11.333, 5.1, CARD_BG, ACCENT_TEAL)

    tb = slide1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "HACKATHON TRACK: AIML-08"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_TEAL

    p1 = tf.add_paragraph()
    p1.text = "PriceGuard AI 🛡️"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "Centralized Automated System for Price Estimation & Reasonability"
    p2.font.size = Pt(20)
    p2.font.color.rgb = TEXT_MUTED

    p3 = tf.add_paragraph()
    p3.text = "\nSDG Alignment: SDG 1 (No Poverty) | SDG 8 (Decent Work) | SDG 12 (Responsible Consumption)"
    p3.font.size = Pt(13)
    p3.font.color.rgb = ACCENT_CYAN

    slide1.notes_slide.notes_text_frame.text = (
        "Good morning judges! Today we present PriceGuard AI — a centralized automated intelligence platform "
        "designed to monitor market commodities, compute dynamic fair prices, and catch price manipulation before it harms consumers."
    )

    # ==========================================
    # SLIDE 2: Problem Statement
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_background(slide2)
    add_header(slide2, "The Problem: Market Gouging & Lack of Price Transparency")

    # 3 Cards
    cards_data2 = [
        ("🚨 Artificial Price Spikes", "Essential food staples like onions and tomatoes suffer 60%+ artificial price spikes caused by hoarding and supply manipulation."),
        ("👁️ Zero Real-time Visibility", "Neither consumers nor regulatory bodies have automated statistical baselines to determine if a market price is fair."),
        ("📉 Economic Vulnerability", "Uncontrolled price inflation hits lower-income households hardest, directly threatening household food security.")
    ]

    for i, (title, desc) in enumerate(cards_data2):
        left = 0.8 + i * 4.0
        add_card(slide2, left, 1.8, 3.7, 4.8, CARD_BG)
        
        tb = slide2.shapes.add_textbox(Inches(left + 0.2), Inches(2.1), Inches(3.3), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_RED

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_MUTED

    slide2.notes_slide.notes_text_frame.text = (
        "Every season, essential food prices skyrocket overnight. The main issue is artificial hoarding and market opacity. "
        "Consumers don't know what a fair price is, and authorities lack real-time tools to spot market manipulation."
    )

    # ==========================================
    # SLIDE 3: The Solution
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_background(slide3)
    add_header(slide3, "The Solution: PriceGuard AI Platform")

    features = [
        ("📊 Automated Fair Value Estimation", "Calculates dynamic statistical baseline price ranges (Mean ± 1 Standard Deviation) using historical feeds."),
        ("🤖 Dual-Layer Anomaly Detection", "Combines Z-Score thresholding with an Isolation Forest Machine Learning model to isolate price gouging."),
        ("🚥 Reasonability Index (0-100)", "Translates complex mathematical metrics into a clean, color-coded score (Fair, Watch, Alert) for non-technical users."),
        ("🚨 Real-Time Alerts Feed", "Generates priority-ranked warning cards for regulatory authorities whenever suspicious volatility occurs.")
    ]

    for i, (ft_title, ft_desc) in enumerate(features):
        row = i // 2
        col = i % 2
        left = 0.8 + col * 5.9
        top = 1.8 + row * 2.6

        add_card(slide3, left, top, 5.6, 2.3, CARD_BG, ACCENT_TEAL if i==1 else None)

        tb = slide3.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(5.2), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = ft_title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_TEAL

        p2 = tf.add_paragraph()
        p2.text = "\n" + ft_desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_LIGHT

    slide3.notes_slide.notes_text_frame.text = (
        "PriceGuard AI acts as an automated guardian for market prices. It processes price feeds, runs ML algorithms, "
        "computes a 0-100 Reasonability Score, and immediately flags price manipulation."
    )

    # ==========================================
    # SLIDE 4: SDG Alignment
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_background(slide4)
    add_header(slide4, "UN Sustainable Development Goals (SDGs)")

    sdgs = [
        ("SDG 1: No Poverty", RGBColor(229, 36, 59), "Protects low-income households from predatory price gouging on daily essential food staples."),
        ("SDG 8: Decent Work & Economic Growth", RGBColor(162, 25, 66), "Fosters transparent, fair commercial market practices and protects legitimate sellers."),
        ("SDG 12: Responsible Consumption", RGBColor(191, 139, 46), "Enables data-backed price accountability across the entire supply chain from farm to retail.")
    ]

    for i, (sdg_title, color, sdg_desc) in enumerate(sdgs):
        top = 1.8 + i * 1.7
        add_card(slide4, 0.8, top, 11.7, 1.4, CARD_BG, color)

        tb = slide4.shapes.add_textbox(Inches(1.1), Inches(top + 0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = sdg_title
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = sdg_desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_LIGHT

    slide4.notes_slide.notes_text_frame.text = (
        "PriceGuard AI directly aligns with 3 UN SDGs: SDG 1 by protecting household budgets, SDG 8 by enforcing fair market practices, "
        "and SDG 12 by bringing end-to-end supply chain transparency."
    )

    # ==========================================
    # SLIDE 5: System Architecture & Workflow
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_background(slide5)
    add_header(slide5, "System Architecture & Data Pipeline")

    steps = [
        ("1. Data Ingestion", "Historical & current market feeds (eNAM / AGMARKNET)"),
        ("2. Stat Baseline", "Calculate Mean, Std Dev & Fair Range Bounds"),
        ("3. Isolation Forest ML", "Multi-feature outlier detection & Z-Score validation"),
        ("4. Flask REST API", "Serves JSON payloads to client interfaces"),
        ("5. Interactive UI", "Dashboard rendering with Chart.js & alert cards")
    ]

    for i, (step_t, step_d) in enumerate(steps):
        left = 0.8 + i * 2.38
        add_card(slide5, left, 2.5, 2.2, 3.5, CARD_BG, ACCENT_CYAN if i==2 else None)

        tb = slide5.shapes.add_textbox(Inches(left + 0.15), Inches(2.7), Inches(1.9), Inches(3.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = step_t
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_CYAN

        p2 = tf.add_paragraph()
        p2.text = "\n" + step_d
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED

    slide5.notes_slide.notes_text_frame.text = (
        "Our data flow starts from market feeds, computes statistical baselines, feeds normalized feature vectors into an Isolation Forest model, "
        "and serves the payload via a Flask API to our interactive dashboard."
    )

    # ==========================================
    # SLIDE 6: Machine Learning Engine
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_background(slide6)
    add_header(slide6, "Machine Learning & Analytics Engine")

    ml_cards = [
        ("📐 Statistical Fair Range", "Uses Mean ± 1 Standard Deviation to define the equilibrium price band where ~68% of normal historical prices reside."),
        ("📊 Z-Score Thresholding", "Z = (Price - Mean) / StdDev. Measures how many standard deviations a price strays. Values |Z| > 2 flag statistical anomalies."),
        ("🌲 Isolation Forest ML", "Unsupervised tree ensemble trained on [price_ratio, volatility_cv, total_change]. Isolates abnormal pricing patterns independently of linear rules."),
        ("📈 Linear Regression Trend", "Fits slope line across historical sequences to differentiate gradual inflation from sudden temporary price spikes.")
    ]

    for i, (m_title, m_desc) in enumerate(ml_cards):
        row = i // 2
        col = i % 2
        left = 0.8 + col * 5.9
        top = 1.8 + row * 2.6

        add_card(slide6, left, top, 5.6, 2.3, CARD_BG)

        tb = slide6.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(5.2), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = m_title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_TEAL

        p2 = tf.add_paragraph()
        p2.text = "\n" + m_desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_LIGHT

    slide6.notes_slide.notes_text_frame.text = (
        "We use a multi-layered analytical engine: Z-Scores for statistical distance, Linear Regression for long-term trends, "
        "and Isolation Forest Machine Learning to detect non-linear pricing anomalies."
    )

    # ==========================================
    # SLIDE 7: Dashboard Features
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_background(slide7)
    add_header(slide7, "Product Showcase: Live Dashboard UI")

    add_card(slide7, 0.8, 1.8, 5.6, 4.8, CARD_BG)
    tb_l = slide7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "🖥️ Key UI Components"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    points_l = [
        "Status Overview Cards: Overall market health, average reasonability, active alerts.",
        "Commodity Grid Cards: Real-time price, fair min/max range, status badges (Fair / Watch / Alert).",
        "Visual Reasonability Bar: Color-coded progress bar (0-100%) indicating price fairness.",
        "Historical Trend Chart: Chart.js powered 12-month graph with dashed fair boundaries."
    ]
    for pt in points_l:
        p = tf_l.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT

    add_card(slide7, 6.7, 1.8, 5.8, 4.8, CARD_BG)
    tb_r = slide7.shapes.add_textbox(Inches(6.9), Inches(2.0), Inches(5.4), Inches(4.4))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "🚨 Anomaly Feed & Alerts"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED

    points_r = [
        "High Severity Alerts: Triggered when Z > 2.0 or score drops below 40%.",
        "Medium Severity Warnings: Issued for sustained inflation above fair ceilings.",
        "Low Severity Monitoring: Logs minor seasonal adjustments.",
        "Full Context Details: Describes percentage surge, statistical metrics, and trigger source."
    ]
    for pt in points_r:
        p = tf_r.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT

    slide7.notes_slide.notes_text_frame.text = (
        "Our web interface translates complex data into simple visual cards, status lights, trend graphs, and prioritized alerts."
    )

    # ==========================================
    # SLIDE 8: Tech Stack & Executable
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_background(slide8)
    add_header(slide8, "Technical Stack & Standalone Executable")

    stack_cards = [
        ("🎨 Frontend", "HTML5, Glassmorphic CSS3, ES6+ JS, Chart.js"),
        ("⚡ Backend Server", "Python 3.13, Flask REST API"),
        ("🧠 ML & Stats", "NumPy, Scikit-learn (IsolationForest, LinearRegression)"),
        ("📦 Distribution", "Compiled Standalone Windows Executable (PriceGuardAI.exe)")
    ]

    for i, (st_t, st_d) in enumerate(stack_cards):
        left = 0.8 + i * 2.95
        add_card(slide8, left, 1.8, 2.7, 4.8, CARD_BG, ACCENT_TEAL if i==3 else None)

        tb = slide8.shapes.add_textbox(Inches(left + 0.15), Inches(2.0), Inches(2.4), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = st_t
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_TEAL

        p2 = tf.add_paragraph()
        p2.text = "\n" + st_d
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_LIGHT

    slide8.notes_slide.notes_text_frame.text = (
        "We built our system using Python, Flask, Scikit-Learn, and Vanilla JS. "
        "For seamless deployment, we compiled everything into a standalone PriceGuardAI.exe file that runs with zero dependencies."
    )

    # ==========================================
    # SLIDE 9: Roadmap & Scaling
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_background(slide9)
    add_header(slide9, "Future Roadmap & Scalability")

    phases = [
        ("Phase 1: Current Baseline", "Functional ML model, REST API, responsive web dashboard, and standalone desktop package."),
        ("Phase 2: Live Mandi API Integration", "Direct API integration with national agriculture feeds (eNAM, AGMARKNET) for real-time market data."),
        ("Phase 3: Multi-Regional Geo-Pricing", "District-level fair estimation considering transportation, regional tariffs, and local supply constraints."),
        ("Phase 4: Predictive Forecasting & Push Alerts", "LSTM deep learning models for advance price forecasting and automated SMS/WhatsApp alerts for market inspectors.")
    ]

    for i, (ph_t, ph_d) in enumerate(phases):
        top = 1.8 + i * 1.3
        add_card(slide9, 0.8, top, 11.7, 1.1, CARD_BG)

        tb = slide9.shapes.add_textbox(Inches(1.0), Inches(top + 0.1), Inches(11.3), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = ph_t
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_CYAN

        p2 = tf.add_paragraph()
        p2.text = ph_d
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_LIGHT

    slide9.notes_slide.notes_text_frame.text = (
        "Our roadmap includes integrating live government APIs, adding district-level geo-pricing, and deploying predictive LSTM models with automated SMS alerts."
    )

    # ==========================================
    # SLIDE 10: Conclusion & Q&A
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_background(slide10)

    add_card(slide10, 1.0, 1.2, 11.333, 5.1, CARD_BG, ACCENT_TEAL)

    tb = slide10.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "PRICEGUARD AI 🛡️"
    p0.font.size = Pt(36)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_TEAL

    p1 = tf.add_paragraph()
    p1.text = "Fair Markets Through Machine Learning"
    p1.font.size = Pt(22)
    p1.font.color.rgb = TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "\n• Automated Data-Driven Fair Estimation\n• Dual-Layer Machine Learning Outlier Detection\n• Zero-Setup Standalone Deployment (.exe)"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED

    p3 = tf.add_paragraph()
    p3.text = "\nThank You! Open for Questions & Feedback."
    p3.font.size = Pt(18)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_CYAN

    slide10.notes_slide.notes_text_frame.text = (
        "PriceGuard AI brings automation, transparency, and machine learning to essential commodity pricing. "
        "Thank you for your time, and we look forward to your questions!"
    )

    # Save presentation
    output_path = os.path.join(os.getcwd(), "PriceGuard_AI_Presentation.pptx")
    prs.save(output_path)
    print(f"✅ Presentation saved to: {output_path}")

if __name__ == "__main__":
    import sys, io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    create_presentation()
